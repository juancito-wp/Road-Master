"""Construcción de reportes y facturas en PDF y Excel (reportlab + openpyxl)."""

import io
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..common import to_float

PROYECTO = 'ROAD MASTER'
ROJO = colors.HexColor('#b91c1c')
GRIS = colors.HexColor('#e2e8f0')


def _dinero(valor) -> str:
    return f'${to_float(valor):,.0f}'


def _ruta_logo() -> Path:
    return Path(__file__).resolve().parent.parent / 'assets' / 'logo_road_master.png'


def _encabezado(estilos: dict, subtitulo: str) -> list:
    """Encabezado del PDF: logo de Road Master junto al título del documento."""
    titulos = [Paragraph(PROYECTO, estilos['titulo']), Paragraph(subtitulo, estilos['subtitulo'])]
    if not _ruta_logo().exists():
        return titulos
    logo = PdfImage(str(_ruta_logo()), width=2.2 * cm, height=2.2 * cm)
    tabla = Table([[logo, titulos]], colWidths=[2.8 * cm, 15.6 * cm])
    tabla.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (0, 0), 0),
        ('RIGHTPADDING', (0, 0), (0, 0), 0),
        ('LEFTPADDING', (1, 0), (1, 0), 10),
    ]))
    return [tabla]


def _logo_excel(hoja) -> int:
    """Inserta el logo en la esquina superior del Excel y devuelve las filas que ocupa."""
    if not _ruta_logo().exists():
        return 0
    imagen = ExcelImage(str(_ruta_logo()))
    imagen.width, imagen.height = 95, 95
    hoja.add_image(imagen, 'A1')
    for fila in range(1, 6):
        hoja.row_dimensions[fila].height = 20
    return 5


def _estilos() -> dict:
    base = getSampleStyleSheet()
    return {
        'titulo': ParagraphStyle('titulo', parent=base['Title'], fontSize=20, textColor=ROJO, spaceAfter=4),
        'subtitulo': ParagraphStyle('subtitulo', parent=base['Normal'], fontSize=11, textColor=colors.HexColor('#334155')),
        'seccion': ParagraphStyle('seccion', parent=base['Normal'], fontSize=10, leading=14),
        'pie': ParagraphStyle('pie', parent=base['Normal'], fontSize=8, textColor=colors.HexColor('#64748b')),
    }


def _tabla(filas: list[list[str]], anchos: list[float]) -> Table:
    tabla = Table(filas, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), ROJO),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    return tabla


def _documento(buffer: io.BytesIO, titulo: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        buffer, pagesize=letter, title=titulo,
        leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm,
    )


def construir_pdf_factura(factura, cliente, detalles) -> bytes:
    """Genera el PDF de una factura de venta."""
    estilos = _estilos()
    buffer = io.BytesIO()
    doc = _documento(buffer, f'Factura {factura.numero}')
    generado = datetime.now().strftime('%Y-%m-%d %H:%M')

    cliente_texto = (
        f"<b>Cliente:</b> {cliente.nombre} {cliente.apellido}<br/>"
        f"<b>Documento:</b> {cliente.numero_documento} &nbsp;&nbsp; <b>Teléfono:</b> {cliente.telefono}<br/>"
        f"<b>Correo:</b> {cliente.email}<br/>"
        f"<b>Dirección:</b> {cliente.direccion}"
        if cliente else '<b>Cliente:</b> No disponible'
    )

    elementos = [
        *_encabezado(estilos, 'FACTURA DE VENTA'),
        Spacer(1, 10),
        _tabla([
            ['Número de factura', factura.numero],
            ['Fecha de emisión', factura.fecha.strftime('%Y-%m-%d %H:%M')],
            ['Estado', factura.estado.upper()],
            ['Venta asociada', f'#{factura.venta_id}'],
        ], [4.5 * cm, 12 * cm]),
        Spacer(1, 12),
        Paragraph(cliente_texto, estilos['seccion']),
        Spacer(1, 12),
    ]

    filas = [['Descripción', 'Cantidad', 'Precio unitario', 'Subtotal']]
    for detalle in detalles:
        filas.append([
            detalle.descripcion[:70],
            f'{to_float(detalle.cantidad):.0f}',
            _dinero(detalle.precio_unitario),
            _dinero(detalle.subtotal),
        ])
    filas.append(['', '', 'Subtotal', _dinero(factura.subtotal)])
    filas.append(['', '', 'Impuestos', _dinero(factura.impuesto)])
    filas.append(['', '', 'TOTAL', _dinero(factura.total)])
    tabla = _tabla(filas, [8.5 * cm, 2.2 * cm, 3 * cm, 2.8 * cm])
    tabla.setStyle(TableStyle([
        ('FONTNAME', (2, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (2, -1), (-1, -1), GRIS),
    ]))
    elementos += [tabla, Spacer(1, 16),
                  Paragraph(f'Documento generado automáticamente por {PROYECTO} el {generado}.', estilos['pie'])]

    doc.build(elementos)
    return buffer.getvalue()


def construir_pdf_reporte_diario(fecha, ventas: list[dict], resumen: dict) -> bytes:
    """Genera el PDF del reporte diario de ventas."""
    estilos = _estilos()
    buffer = io.BytesIO()
    doc = _documento(buffer, f'Reporte diario de ventas {fecha}')
    generado = datetime.now().strftime('%Y-%m-%d %H:%M')

    filas = [['Fecha', 'N° venta', 'Cliente', 'Producto / Servicio', 'Cant.', 'Valor unit.', 'Total venta', 'Estado']]
    for venta in ventas:
        detalles = venta['detalles'] or [{'descripcion': 'Sin detalle', 'cantidad': 0, 'precioUnitario': 0}]
        for indice, detalle in enumerate(detalles):
            filas.append([
                (venta['fecha'] or '')[:10] if indice == 0 else '',
                f"#{venta['id']}" if indice == 0 else '',
                ((venta['cliente'] or {}).get('nombre', '') + ' ' + (venta['cliente'] or {}).get('apellido', '')) if indice == 0 else '',
                detalle['descripcion'][:45],
                f"{detalle['cantidad']:.0f}",
                _dinero(detalle['precioUnitario']),
                _dinero(venta['total']) if indice == 0 else '',
                venta['estado'].upper() if indice == 0 else '',
            ])

    elementos = [
        *_encabezado(estilos, 'REPORTE DIARIO DE VENTAS'),
        Spacer(1, 10),
        _tabla([
            ['Fecha del reporte', fecha],
            ['Ventas registradas', str(resumen['cantidadVentas'])],
            ['Total vendido', _dinero(resumen['totalVendido'])],
            ['Total impuestos', _dinero(resumen['totalImpuestos'])],
        ], [4.5 * cm, 12 * cm]),
        Spacer(1, 12),
        _tabla(filas, [1.9 * cm, 1.8 * cm, 3.4 * cm, 4.3 * cm, 1.3 * cm, 2 * cm, 2 * cm, 1.8 * cm]),
        Spacer(1, 16),
        Paragraph(f'Informe generado por {PROYECTO} el {generado} — usuario del sistema.', estilos['pie']),
    ]
    doc.build(elementos)
    return buffer.getvalue()


def construir_excel_reporte_diario(fecha, ventas: list[dict], resumen: dict) -> bytes:
    """Genera el Excel (.xlsx) del reporte diario de ventas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Reporte diario'

    titulo = Font(bold=True, size=14, color='B91C1C')
    encabezado = Font(bold=True, color='FFFFFF')
    relleno = PatternFill('solid', fgColor='B91C1C')
    borde = Border(*[Side(style='thin', color='CBD5E1')] * 4)

    filas_logo = _logo_excel(hoja)
    for _ in range(filas_logo):
        hoja.append([])
    hoja.append([f'{PROYECTO} - Reporte diario de ventas'])
    hoja.cell(row=filas_logo + 1, column=1).font = titulo
    hoja.append([f'Fecha del reporte: {fecha}'])
    hoja.append([f"Ventas: {resumen['cantidadVentas']}",
                 f"Total vendido: {to_float(resumen['totalVendido'])}",
                 f"Impuestos: {to_float(resumen['totalImpuestos'])}"])
    hoja.append([])

    columnas = ['Fecha', 'N° venta', 'Cliente', 'Producto / Servicio', 'Tipo', 'Cantidad',
                'Valor unitario', 'Descuento', 'Subtotal línea', 'Total venta', 'Estado']
    hoja.append(columnas)
    fila_encabezado = hoja.max_row
    for celda in hoja[fila_encabezado]:
        celda.font = encabezado
        celda.fill = relleno
        celda.alignment = Alignment(horizontal='center', vertical='center')

    for venta in ventas:
        cliente = venta['cliente'] or {}
        for detalle in venta['detalles'] or []:
            hoja.append([
                (venta['fecha'] or '')[:10],
                venta['id'],
                f"{cliente.get('nombre', '')} {cliente.get('apellido', '')}".strip(),
                detalle['descripcion'],
                detalle['tipo'],
                to_float(detalle['cantidad']),
                to_float(detalle['precioUnitario']),
                to_float(detalle['descuento']),
                to_float(detalle['subtotal']),
                to_float(venta['total']),
                venta['estado'],
            ])

    for fila in hoja.iter_rows(min_row=fila_encabezado + 1, max_row=hoja.max_row):
        for celda in fila:
            celda.border = borde
        for indice in (6, 7, 8, 9):
            fila[indice].number_format = '"$"#,##0.00'

    hoja.freeze_panes = f'A{fila_encabezado + 1}'
    hoja.auto_filter.ref = f'A{fila_encabezado}:K{max(hoja.max_row, fila_encabezado)}'
    anchos = [11, 9, 24, 34, 11, 9, 15, 12, 14, 13, 11]
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.append([])
    hoja.append(['', '', '', '', '', '', '', 'TOTAL GENERAL', Decimal(to_float(resumen['totalVendido']))])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
